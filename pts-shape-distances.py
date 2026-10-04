"""
TODO:
    Research strategies on dynamically building road shapefiles for unfinished sections
"""

# -----------------------------------------------------------------------------
# Authors: Cirenio Sanchez, Dr.DingXin Cheng
# Date: 2025-09-19
# -----------------------------------------------------------------------------

import os
import re
import math
import shutil
import threading
import datetime as dt
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
from copy import copy

import pandas as pd
import geopandas as gpd
from shapely.geometry import LineString, MultiLineString, Point
from decimal import Decimal, ROUND_HALF_UP
import pyproj  # for UTM-based inter-point distance

# ---------- constants / defaults ----------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_QGIS_PY_BAT = r"C:\Program Files\QGIS 3.40.14\bin\python-qgis-ltr.bat"
DEFAULT_QGIS_PROJECT = os.path.join(SCRIPT_DIR, "shortest-dist-gis-map.qgz")
DEFAULT_QUERY_SCRIPT = os.path.join(SCRIPT_DIR, "query.py")  # sanitizer script
DEFAULT_GPS_BASED_LENGTH = True

# Template path (relative to project/script folder)
TEMPLATE_EXCEL = os.path.join(SCRIPT_DIR, "output", "sample_output", "SSInspectionImport.xlsx")

# ---------- helpers ----------
def utm_epsg_for_roads(lon: float, lat: float) -> str:
    zone = int(math.floor((lon + 180) / 6) + 1)
    return f"EPSG:{32600 + zone if lat >= 0 else 32700 + zone}"


def utm_epsg_for_distance(lat: float, lon: float) -> int:
    zone = int((lon + 180) / 6) + 1
    hemisphere = 32600 if lat >= 0 else 32700
    return hemisphere + zone


def utm_distance_m(lat1, lon1, lat2, lon2) -> float:
    epsg = utm_epsg_for_distance(lat1, lon1)
    proj = pyproj.CRS.from_epsg(epsg)
    transformer = pyproj.Transformer.from_crs("EPSG:4326", proj, always_xy=True)
    x1, y1 = transformer.transform(lon1, lat1)
    x2, y2 = transformer.transform(lon2, lat2)
    return math.dist((x1, y1), (x2, y2))


def utm_distance_ft(lat1, lon1, lat2, lon2) -> float:
    meters = utm_distance_m(lat1, lon1, lat2, lon2)
    return meters * 3.280839895013123


def extract_date_from_point_path(path_str: str) -> str:
    if path_str is None:
        return ""
    s = str(path_str)
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if not m:
        return ""
    y, mo, d = map(int, m.groups())
    return f"{mo:02d}/{d:02d}/{y:04d}"


def split_distress(col: str):
    parts = col.split("-")
    return (parts[0], "") if len(parts) == 1 else ("-".join(parts[:-1]), parts[-1])


def round1(x: float) -> float:
    return float(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def round1_or_none(x: float):
    return round1(x) if (x is not None and math.isfinite(x)) else None


def ensure_dir(path: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)


def safe_remove(path: str):
    try:
        if os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def write_excel(df: pd.DataFrame, path: str) -> str:
    import os
    import shutil
    import pandas as pd
    from copy import copy
    from openpyxl import load_workbook

    script_dir = os.path.dirname(os.path.abspath(__file__))
    template_excel = os.path.join(script_dir, "output", "sample_output", "SSInspectionImport.xlsx")

    if not os.path.isfile(template_excel):
        raise FileNotFoundError(f"Excel template not found: {template_excel}")

    os.makedirs(os.path.dirname(path), exist_ok=True)

    # 1) Copy template -> destination (preserves template formatting)
    shutil.copyfile(template_excel, path)

    # 2) Load copied workbook
    wb = load_workbook(path)

    # 3) Pick sheet
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

    df_to_write = df.copy()

    HEADER_ROW = 1
    STYLE_ROW = 2

    if ws.max_row < STYLE_ROW:
        raise ValueError("Template must include row 2 as a formatted blank data row.")

    # Build header_map from row 1
    header_map = {}
    for col_idx in range(1, ws.max_column + 1):
        v = ws.cell(row=HEADER_ROW, column=col_idx).value
        if v is not None and str(v).strip():
            header_map[str(v).strip()] = col_idx

    if not header_map:
        raise ValueError("Template row 1 header appears empty.")

    def copy_cell_style(src_cell, dst_cell):
        dst_cell._style = copy(src_cell._style)
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.number_format = src_cell.number_format

    def apply_template_style(dst_row: int):
        """
        Data rows:
          - Copy MOST styling from STYLE_ROW (row 2)
          - Copy ALIGNMENT from HEADER_ROW (row 1)
        """
        for col_idx in range(1, ws.max_column):
            header_cell = ws.cell(row=HEADER_ROW, column=col_idx)
            style_cell = ws.cell(row=STYLE_ROW, column=col_idx)
            dst_cell = ws.cell(row=dst_row, column=col_idx)

            # base style from row 2
            copy_cell_style(style_cell, dst_cell)

            # alignment from row 1
            dst_cell.alignment = copy(header_cell.alignment)

            # ensure TEXT category
            dst_cell.number_format = "@"

    # Clear old data rows but keep header + style rows
    if ws.max_row > STYLE_ROW:
        ws.delete_rows(STYLE_ROW + 1, ws.max_row - STYLE_ROW)

    start_row = STYLE_ROW  # first record overwrites row 2

    for i, (_, row) in enumerate(df_to_write.iterrows()):
        excel_row = start_row + i

        if excel_row == STYLE_ROW:
            # overwrite style row (keep it styled)
            apply_template_style(excel_row)
        else:
            # insert new row and style it
            ws.insert_rows(excel_row)
            apply_template_style(excel_row)

        # Write values into the proper template columns
        for col_name, val in row.items():
            if col_name not in header_map:
                continue
            col_idx = header_map[col_name]
            cell = ws.cell(row=excel_row, column=col_idx)

            # always write strings to prevent Excel auto-typing
            if str(val).strip():
                cell.value = str(val)
                if col_idx in (header_map.get("Comments"),):
                    # ensure TEXT category
                    cell.number_format = "@"

    wb.save(path)
    return path


def pick_date_token(frames_csv_path: str, inspection_dates: list[str]) -> str:
    for val in inspection_dates:
        if isinstance(val, str) and val.strip():
            try:
                m, d, y = val.split("/")
                return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
            except Exception:
                pass
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", frames_csv_path)
    if m:
        return m.group(0)
    return dt.date.today().isoformat()


def extract_video_name(point_id: str) -> str:
    if not isinstance(point_id, str):
        return ""
    m = re.search(r"([^/\\]+\.mp4)", point_id, flags=re.IGNORECASE)
    return m.group(1).lower() if m else ""


def label_distress(code: str) -> str:
    return DISTRESS_LABEL.get(code, code)


def build_global_endpoints_by_roadname(roads_gdf: gpd.GeoDataFrame, road_name_field: str):
    vertex_counts = {}   # (name_norm, (x, y)) -> count
    vertex_points = {}   # (name_norm, (x, y)) -> Point

    def iter_lines(geom):
        if geom is None or geom.is_empty:
            return
        if isinstance(geom, LineString):
            yield geom
        elif isinstance(geom, MultiLineString):
            for line in geom.geoms:
                if isinstance(line, LineString):
                    yield line

    for _, row in roads_gdf.iterrows():
        geom = row.geometry
        name_val = row.get(road_name_field, "")
        name_norm = str(name_val).strip().upper()
        if not name_norm:
            continue

        for line in iter_lines(geom):
            coords = list(line.coords)
            if len(coords) < 2:
                continue
            for coord in (coords[0], coords[-1]):
                key = (name_norm, (coord[0], coord[1]))
                vertex_counts[key] = vertex_counts.get(key, 0) + 1
                if key not in vertex_points:
                    vertex_points[key] = Point(coord[0], coord[1])

    endpoints_by_name = {}
    for (name_norm, coord), count in vertex_counts.items():
        if count == 1:
            pt = vertex_points[(name_norm, coord)]
            endpoints_by_name.setdefault(name_norm, []).append(pt)

    return endpoints_by_name


def run_subprocess_and_stream(cmd_list, cwd, log_box, btn=None, done_label=None):
    def _runner():
        try:
            p = subprocess.Popen(
                cmd_list,
                cwd=cwd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                shell=False
            )
            for line in p.stdout:
                append_log(log_box, line)
            p.wait()
            rc = p.returncode
            append_log(log_box, f"\n[Sanitizer finished with exit code {rc}]\n")
        except FileNotFoundError as e:
            append_log(log_box, f"[ERROR] Executable not found: {e}\n")
        except Exception as e:
            append_log(log_box, f"[ERROR] {e}\n")
        finally:
            if btn is not None:
                btn.config(state="normal", text=(done_label or "Run"))

    threading.Thread(target=_runner, daemon=True).start()


# ---------- parameters ----------
INSPECTION_UNIT = 1
NO_DISTRESSES = "No"
SPECIAL = "No"

FT_TO_M = 0.3048

IMAGE_LENGTH_FT = 12.0
DEFAULT_LANE_WIDTH_FT = 12.0

R1_FEET_PER_LANE = 40.0
R2_FEET_PER_LANE = 18.0

INTERSECTION_ENDPOINT_NEAR_FT = 60.0

drop_exact_duplicates = True
aggregate_sizes = "sum"
min_distress_size = 0.0

# NEW: drop any scaled distress < 1 from outputs
MIN_OUTPUT_DISTRESS_SIZE = 1.0

XY3857_DECIMALS = 3

DISTRESS_LABEL = {
    "Alligator": "Alligator Cracking",
    "Block": "Block Cracking",
    "L-T": "Long. & Trans. Cracking",
    "Patch": "Patch & Util. Cut Patch",
    "Ravel": "Raveling",
    "Weather": "Weathering",
    "Pothole": "Pothole",
}
DISTRESS_COLS = [
    "Alligator-H", "Alligator-L", "Alligator-M",
    "Block-H", "Block-L", "Block-M",
    "L-T-H", "L-T-L", "L-T-M",
    "Patch-H", "Patch-L", "Patch-M",
    "Ravel-H", "Ravel-M",
    "Weather-H", "Weather-L", "Weather-M",
    "Pothole"
]

start_frame_list = []
end_frame_list = []
# ---------- core processing ----------
def run_pipeline(roads_path: str, frames_csv_path: str, include_intersections: bool = False, use_gps_based_length: bool = DEFAULT_GPS_BASED_LENGTH) -> str:
    # ----- load / clean frames -----
    frames_df = pd.read_csv(frames_csv_path, header=0, skipinitialspace=True)
    frames_df.columns = [c.strip() for c in frames_df.columns]
    present_distress_cols = [c for c in DISTRESS_COLS if c in frames_df.columns]

    numeric_cols = present_distress_cols + ["Latitude", "Longitude", "Width", "Length"]
    for col in numeric_cols:
        if col in frames_df.columns:
            frames_df[col] = pd.to_numeric(frames_df[col], errors="coerce")

    if "Latitude" not in frames_df.columns or "Longitude" not in frames_df.columns:
        raise ValueError("Input CSV must contain Latitude and Longitude columns.")

    frames_df = frames_df[
        frames_df["Latitude"].between(-90, 90) &
        frames_df["Longitude"].between(-180, 180)
    ].dropna(subset=["Latitude", "Longitude"]).copy()

    # Resolve a path column for date + video
    raw_path_col = None
    for field_name in ("PointID", "Image-Name", "ImageName", "Image_Path", "FrameID"):
        if field_name in frames_df.columns:
            raw_path_col = field_name
            break

    if raw_path_col is None:
        frames_df["PointID"] = [f"idx-{i}" for i in frames_df.index]
        raw_path_col = "PointID"

    frames_df["VideoName"] = frames_df[raw_path_col].apply(extract_video_name)

    # ----- GPS-based length between consecutive points per video -----
    length_from_gps_ft = pd.Series(index=frames_df.index, dtype=float)

    def _input_length_or_default(idx) -> float:
        if "Length" in frames_df.columns:
            raw_len = frames_df.at[idx, "Length"]
            if pd.notna(raw_len) and float(raw_len) > 0:
                return float(raw_len)
        return float(IMAGE_LENGTH_FT)

    for vid, group in frames_df.groupby("VideoName", dropna=False):
        idxs_sorted = sorted(list(group.index))
        start_frame_list.append(idxs_sorted[0])
        for i, idx in enumerate(idxs_sorted):
            if use_gps_based_length and i < len(idxs_sorted) - 1:
                lat1 = frames_df.at[idx, "Latitude"]
                lon1 = frames_df.at[idx, "Longitude"]
                idx2 = idxs_sorted[i + 1]
                lat2 = frames_df.at[idx2, "Latitude"]
                lon2 = frames_df.at[idx2, "Longitude"]
                try:
                    d_ft = utm_distance_ft(float(lat1), float(lon1), float(lat2), float(lon2))
                except Exception:
                    d_ft = _input_length_or_default(idx)
            else:
                if i == len(idxs_sorted) - 1:
                    end_frame_list.append(idx)
                d_ft = _input_length_or_default(idx)

            length_from_gps_ft.at[idx] = float(d_ft)

    # ----- build point GeoDataFrames -----
    points_wgs84 = gpd.GeoDataFrame(
        frames_df,
        geometry=gpd.points_from_xy(frames_df["Longitude"], frames_df["Latitude"]),
        crs="EPSG:4326"
    )
    points_webmercator = points_wgs84.to_crs("EPSG:3857")

    # ----- load roads -----
    roads_gdf = gpd.read_file(roads_path)
    if roads_gdf is None or roads_gdf.empty:
        raise ValueError("Roads shapefile is empty.")
    if roads_gdf.crs is None:
        raise ValueError("Roads shapefile has no CRS.")

    road_name_field = "RoadName" if "RoadName" in roads_gdf.columns else (
        "StreetID" if "StreetID" in roads_gdf.columns else None
    )
    if road_name_field is None:
        raise ValueError("Roads must contain 'RoadName' or 'StreetID'.")
    road_id_field = "StreetID" if "StreetID" in roads_gdf.columns else road_name_field

    roads_gdf["lanes_num"] = (
        pd.to_numeric(roads_gdf.get("Lanes", 1), errors="coerce")
        .fillna(1)
        .clip(lower=1)
        .astype(float)
    )

    mean_lon = float(points_wgs84.geometry.x.mean())
    mean_lat = float(points_wgs84.geometry.y.mean())
    epsg_utm_roads = utm_epsg_for_roads(mean_lon, mean_lat)

    points_projected = points_wgs84.to_crs(epsg_utm_roads)
    roads_projected = roads_gdf.to_crs(epsg_utm_roads)

    road_geometries = list(roads_projected.geometry.values)
    road_names = list(roads_projected[road_name_field].astype(str).values)
    road_ids = list(roads_projected[road_id_field].astype(str).values)
    lanes_per_road = list(roads_projected["lanes_num"].astype(float).values)

    if "ShapeKey" in roads_projected.columns:
        corridor_ids = list(roads_projected["ShapeKey"].astype(str).values)
    elif "SectionID" in roads_projected.columns:
        corridor_ids = list(roads_projected["SectionID"].astype(str).values)
    else:
        corridor_ids = road_names

    section_ids = (
        list(roads_projected["SectionID"].astype(str).values)
        if "SectionID" in roads_projected.columns
        else [""] * len(roads_projected)
    )

    endpoints_by_roadname = build_global_endpoints_by_roadname(roads_projected, road_name_field)

    def endpoint_distance_stats_ft(point_geom: Point, road_name: str):
        norm_name = str(road_name).strip().upper()
        endpoints = endpoints_by_roadname.get(norm_name)
        if not endpoints:
            return None, None
        distances_m = [point_geom.distance(ep) for ep in endpoints]
        return min(distances_m) / FT_TO_M, max(distances_m) / FT_TO_M

    # -------------------------------------------------------------------------
    # PASS 1: classify each point + choose main road
    # -------------------------------------------------------------------------
    point_status_by_idx = {}
    chosen_by_idx = {}
    endpoint_override_reason_by_idx = {}

    for point_index, point_row in points_projected.iterrows():
        point_geom = point_row.geometry

        distance_list = [
            (road_index, point_geom.distance(road_geom))
            for road_index, road_geom in enumerate(road_geometries)
        ]

        if not distance_list:
            point_status_by_idx[point_index] = "Not found"
            chosen_by_idx[point_index] = {
                "chosen_road_id": "UNMATCHED",
                "chosen_road_name": "UNMATCHED",
                "chosen_section_id": "",
                "chosen_corridor_id": "",
                "nearest_dist_ft": None,
                "second_dist_ft": None,
                "r1_name": "",
                "r2_name": "",
                "lanes_nearest": 1.0,
                "lanes_second": 1.0,
                "r1_endpoint_count": 0,
                "r2_endpoint_count": 0,
                "r1_min_endpoint_dist_ft": None,
                "r2_min_endpoint_dist_ft": None,
                "r1_max_endpoint_dist_ft": None,
                "r2_max_endpoint_dist_ft": None,
            }
            endpoint_override_reason_by_idx[point_index] = "Not_found"
            continue

        distance_list.sort(key=lambda t: t[1])
        nearest_index, nearest_dist_m = distance_list[0]
        second_index, second_dist_m = (distance_list[1] if len(distance_list) > 1 else (None, float("inf")))

        nearest_dist_ft = nearest_dist_m / FT_TO_M
        second_dist_ft = second_dist_m / FT_TO_M

        lanes_nearest = lanes_per_road[nearest_index] if nearest_index is not None else 1.0
        lanes_second = lanes_per_road[second_index] if second_index is not None else lanes_nearest

        nearest_threshold_ft = R1_FEET_PER_LANE * lanes_nearest
        second_threshold_ft = R2_FEET_PER_LANE * lanes_second

        r1_name = road_names[nearest_index]
        r1_min_endpoint_dist_ft, r1_max_endpoint_dist_ft = endpoint_distance_stats_ft(point_geom, r1_name)
        r1_norm = r1_name.strip().upper()
        r1_endpoint_count = len(endpoints_by_roadname.get(r1_norm, []))

        if second_index is not None:
            r2_name = road_names[second_index]
            r2_min_endpoint_dist_ft, r2_max_endpoint_dist_ft = endpoint_distance_stats_ft(point_geom, r2_name)
            r2_norm = r2_name.strip().upper()
            r2_endpoint_count = len(endpoints_by_roadname.get(r2_norm, []))
        else:
            r2_name = ""
            r2_min_endpoint_dist_ft = None
            r2_max_endpoint_dist_ft = None
            r2_endpoint_count = 0

        is_same_corridor = False
        if second_index is not None:
            is_same_corridor = (
                corridor_ids[nearest_index] == corridor_ids[second_index] or
                road_names[nearest_index].strip().upper() == road_names[second_index].strip().upper()
            )

        if nearest_dist_ft > nearest_threshold_ft:
            point_status = "Not found"
            chosen_index = None
            endpoint_reason = "Not_found"
        elif (second_index is not None) and (second_dist_ft < second_threshold_ft) and (not is_same_corridor):
            point_status = "Intersection"
            chosen_index = nearest_index
            endpoint_reason = "Intersection_initial_R1"
        else:
            point_status = "Normal"
            chosen_index = nearest_index
            endpoint_reason = "Not_intersection"

        if point_status == "Intersection" and second_index is not None:
            chosen_index = nearest_index
            endpoint_reason = "No_endpoint_override_keep_R1"

            if (r1_min_endpoint_dist_ft is not None) and (r2_min_endpoint_dist_ft is not None):
                if (r1_min_endpoint_dist_ft <= INTERSECTION_ENDPOINT_NEAR_FT) and (r2_min_endpoint_dist_ft > INTERSECTION_ENDPOINT_NEAR_FT):
                    chosen_index = second_index
                    endpoint_reason = "Endpoint_override_switch_to_R2_far_endpoints"
                elif (r2_min_endpoint_dist_ft <= INTERSECTION_ENDPOINT_NEAR_FT) and (r1_min_endpoint_dist_ft > INTERSECTION_ENDPOINT_NEAR_FT):
                    chosen_index = nearest_index
                    endpoint_reason = "Endpoint_override_keep_R1_far_endpoints"
                else:
                    if (r1_max_endpoint_dist_ft is not None) and (r2_max_endpoint_dist_ft is not None):
                        if r2_max_endpoint_dist_ft > r1_max_endpoint_dist_ft:
                            chosen_index = second_index
                            endpoint_reason = "Endpoint_override_choose_R2_farthest_endpoint"
                        else:
                            chosen_index = nearest_index
                            endpoint_reason = "Endpoint_override_choose_R1_farthest_endpoint"

        if chosen_index is None:
            chosen_road_name = "UNMATCHED"
            chosen_road_id = "UNMATCHED"
            chosen_section_id = ""
            chosen_corridor_id = ""
        else:
            chosen_road_name = road_names[chosen_index]
            chosen_road_id = road_ids[chosen_index]
            chosen_section_id = section_ids[chosen_index] if section_ids else ""
            chosen_corridor_id = corridor_ids[chosen_index]

        point_status_by_idx[point_index] = point_status
        endpoint_override_reason_by_idx[point_index] = endpoint_reason
        chosen_by_idx[point_index] = {
            "chosen_road_id": chosen_road_id,
            "chosen_road_name": chosen_road_name,
            "chosen_section_id": chosen_section_id,
            "chosen_corridor_id": chosen_corridor_id,
            "nearest_dist_ft": nearest_dist_ft,
            "second_dist_ft": second_dist_ft,
            "r1_name": r1_name,
            "r2_name": r2_name,
            "lanes_nearest": lanes_nearest,
            "lanes_second": lanes_second,
            "r1_endpoint_count": r1_endpoint_count,
            "r2_endpoint_count": r2_endpoint_count,
            "r1_min_endpoint_dist_ft": r1_min_endpoint_dist_ft,
            "r2_min_endpoint_dist_ft": r2_min_endpoint_dist_ft,
            "r1_max_endpoint_dist_ft": r1_max_endpoint_dist_ft,
            "r2_max_endpoint_dist_ft": r2_max_endpoint_dist_ft,
        }

    # -------------------------------------------------------------------------
    # PASS 2: effective length logic + intersection reassignment rule
    # -------------------------------------------------------------------------
    effective_length_ft = length_from_gps_ft.copy().astype(float)

    if not include_intersections:
        for vid, group in frames_df.groupby("VideoName", dropna=False):
            idxs_sorted = sorted(list(group.index))
            for i, idx in enumerate(idxs_sorted):
                if point_status_by_idx.get(idx) != "Normal":
                    continue

                prev_idx = idxs_sorted[i - 1] if i > 0 else None
                next_idx = idxs_sorted[i + 1] if i < len(idxs_sorted) - 1 else None

                touches_intersection = (
                    (prev_idx is not None and point_status_by_idx.get(prev_idx) == "Intersection") or
                    (next_idx is not None and point_status_by_idx.get(next_idx) == "Intersection")
                )

                if touches_intersection:
                    effective_length_ft.at[idx] = _input_length_or_default(idx)

    # -------------------------------------------------------------------------
    # Updated input: new Length + scaled distress values (rounded to 0.1)
    # ALSO: remove rows with 0 distresses after scaling
    # -------------------------------------------------------------------------
    updated_frames_df = frames_df.copy()
    if not updated_frames_df.empty:
        if "Length" in updated_frames_df.columns:
            updated_frames_df["Length"] = updated_frames_df["Length"].astype(float)

        for col in present_distress_cols:
            updated_frames_df[col] = pd.to_numeric(updated_frames_df[col], errors="coerce").astype(float)

        for idx in updated_frames_df.index:
            seg_len = float(effective_length_ft.get(idx, IMAGE_LENGTH_FT))
            if not math.isfinite(seg_len) or seg_len <= 0:
                seg_len = float(IMAGE_LENGTH_FT)

            if idx in end_frame_list:
                 factor = 1.0
            else:
                factor = seg_len / float(IMAGE_LENGTH_FT)

            if "Length" in updated_frames_df.columns:
                if idx in start_frame_list:
                    seg_len = factor * frames_df.at[idx, "Length"] 
                    
               
                updated_frames_df.at[idx, "Length"] = float(
                    Decimal(str(seg_len)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                )

            for col in present_distress_cols:
                val = updated_frames_df.at[idx, col]
                if pd.notna(val) and float(val) > 0:
                    scaled = float(val) * factor
                    updated_frames_df.at[idx, col] = float(
                        Decimal(str(scaled)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                    )
                else:
                    # keep clean zeros
                    updated_frames_df.at[idx, col] = 0.0

        # Drop rows where all distress columns are 0 (or missing)
        if present_distress_cols:
            mask_any_distress = (updated_frames_df[present_distress_cols].fillna(0) > 0).any(axis=1)
            updated_frames_df = updated_frames_df.loc[mask_any_distress].copy()

        # SAFETY: remove VideoName from updated input
        if "VideoName" in updated_frames_df.columns:
            updated_frames_df = updated_frames_df.drop(columns=["VideoName"])

    # -------------------------------------------------------------------------
    # PASS 3: build outputs
    # -------------------------------------------------------------------------
    inspection_rows = []
    point_rows_output = []
    intersection_debug_rows = []
    section_totals = {}  # (StreetID, SectionID, InspectionDate) -> {"length":..., "area":...}

    for point_index, point_row in points_projected.iterrows():
        point_geom = point_row.geometry

        raw_path_value = None
        for field_name in ("PointID", "Image-Name", "ImageName", "Image_Path", "FrameID"):
            if field_name in points_wgs84.columns:
                value = points_wgs84.loc[point_index, field_name]
                if isinstance(value, str) and value.strip():
                    raw_path_value = value.strip()
                    break
        if raw_path_value is None:
            raw_path_value = f"idx-{point_index}"

        frame_id = raw_path_value
        inspection_date = extract_date_from_point_path(raw_path_value)

        point_status = point_status_by_idx.get(point_index, "Not found")
        chosen = chosen_by_idx.get(point_index, {})
        chosen_road_id = chosen.get("chosen_road_id", "UNMATCHED")
        chosen_road_name = chosen.get("chosen_road_name", "UNMATCHED")
        chosen_section_id = chosen.get("chosen_section_id", "")
        chosen_corridor_id = chosen.get("chosen_corridor_id", "")

        if point_status == "Intersection" and not include_intersections:
            point_status_for_points = "Intersection"
        elif point_status == "Intersection" and include_intersections:
            point_status_for_points = "Normal"
        else:
            point_status_for_points = point_status

        if "Width" in points_wgs84.columns:
            width_value_raw = points_wgs84.loc[point_index, "Width"]
            width_ft_value = float(width_value_raw) if pd.notna(width_value_raw) and float(width_value_raw) > 0 else DEFAULT_LANE_WIDTH_FT
        else:
            width_ft_value = DEFAULT_LANE_WIDTH_FT

        inspection_length_ft = float(effective_length_ft.get(point_index, IMAGE_LENGTH_FT))
        if not math.isfinite(inspection_length_ft) or inspection_length_ft <= 0:
            inspection_length_ft = float(IMAGE_LENGTH_FT)

        area_sq_ft = float(width_ft_value) * float(inspection_length_ft)

        if (
            point_status != "Not found"
            and not (point_status == "Intersection" and not include_intersections)
            and chosen_road_id != "UNMATCHED"
        ):
            key_sec = (str(chosen_road_id), str(chosen_section_id), inspection_date)
            rec = section_totals.setdefault(key_sec, {"length": 0.0, "area": 0.0})
            rec["length"] += float(inspection_length_ft)
            rec["area"] += float(area_sq_ft)

        point_geom_webmercator = points_webmercator.loc[point_index, "geometry"]
        point_rows_output.append({
            "PointID": frame_id,
            "X_EPSG3857": round(float(point_geom_webmercator.x), XY3857_DECIMALS),
            "Y_EPSG3857": round(float(point_geom_webmercator.y), XY3857_DECIMALS),
            "PolylineID": chosen_corridor_id if point_status_for_points != "Not found" else "",
            "RoadName": chosen_road_name if point_status_for_points != "Not found" else "UNMATCHED",
            "Status": point_status_for_points,
            "R1_ft": round1_or_none(chosen.get("nearest_dist_ft")) if chosen.get("nearest_dist_ft") is not None else None,
            "R2_ft": round1_or_none(chosen.get("second_dist_ft")) if chosen.get("second_dist_ft") is not None else None,
        })

        dist_to_next_ft = float(length_from_gps_ft.get(point_index, IMAGE_LENGTH_FT))
        intersection_debug_rows.append({
            "PointID": frame_id,
            "RawPointPath": raw_path_value,
            "ParsedInspectionDate": inspection_date,
            "IncludeIntersectionsFlag": bool(include_intersections),
            "InitialR1Name": chosen.get("r1_name", ""),
            "InitialR2Name": chosen.get("r2_name", ""),
            "InitialR1Dist_ft": round1_or_none(chosen.get("nearest_dist_ft")) if chosen.get("nearest_dist_ft") is not None else None,
            "InitialR2Dist_ft": round1_or_none(chosen.get("second_dist_ft")) if chosen.get("second_dist_ft") is not None else None,
            "R1_Lanes": chosen.get("lanes_nearest", None),
            "R2_Lanes": chosen.get("lanes_second", None),
            "R1_EndpointCount": chosen.get("r1_endpoint_count", 0),
            "R2_EndpointCount": chosen.get("r2_endpoint_count", 0),
            "Nearest_R1_EndpointDist_ft": round1_or_none(chosen.get("r1_min_endpoint_dist_ft")) if chosen.get("r1_min_endpoint_dist_ft") is not None else None,
            "Nearest_R2_EndpointDist_ft": round1_or_none(chosen.get("r2_min_endpoint_dist_ft")) if chosen.get("r2_min_endpoint_dist_ft") is not None else None,
            "Farthest_R1_EndpointDist_ft": round1_or_none(chosen.get("r1_max_endpoint_dist_ft")) if chosen.get("r1_max_endpoint_dist_ft") is not None else None,
            "Farthest_R2_EndpointDist_ft": round1_or_none(chosen.get("r2_max_endpoint_dist_ft")) if chosen.get("r2_max_endpoint_dist_ft") is not None else None,
            "DistanceToNextPoint_ft": round1_or_none(dist_to_next_ft),
            "EffectiveLength_ft": round1_or_none(inspection_length_ft),
            "EndpointOverrideReason": endpoint_override_reason_by_idx.get(point_index, ""),
            "ChosenMainRoadName": chosen_road_name,
            "ChosenMainRoadID": chosen_road_id,
            "ChosenSectionID": chosen_section_id,
            "Status": point_status,
        })

        # main inspection rows
        if point_status == "Not found":
            continue
        if (point_status == "Intersection") and (not include_intersections):
            continue

        factor_for_size = float(inspection_length_ft) / float(IMAGE_LENGTH_FT)

        for col in present_distress_cols:
            size_value = points_wgs84.loc[point_index, col]
            if pd.isna(size_value) or float(size_value) <= min_distress_size:
                continue

            distress_base, severity = split_distress(col)
            scaled_size = float(size_value) * factor_for_size

            # NEW: drop any output distress < 1 (after scaling)
            if (not math.isfinite(scaled_size)) or (scaled_size < MIN_OUTPUT_DISTRESS_SIZE):
                continue

            inspection_rows.append({
                "FrameID": frame_id,
                "FrameWeight": 1.0,
                "StreetID": chosen_road_id,
                "SectionID": chosen_section_id,
                "InspectionUnit.": INSPECTION_UNIT,
                "InspectionDate": inspection_date,
                "DistressType": label_distress(distress_base),
                "Severity": severity,
                "RawSize": float(scaled_size),
                "InspectionArea": float(area_sq_ft),
                "InspectionLength": float(inspection_length_ft),
                "NoDistresses": NO_DISTRESSES,
                "Special": SPECIAL,
                "Comment": "",
            })

    # ----- aggregate main inspection dataframe -----
    main_df = pd.DataFrame(inspection_rows)
    if not main_df.empty and drop_exact_duplicates:
        main_df = main_df.drop_duplicates()

    if not main_df.empty:
        group_keys = ["StreetID", "SectionID", "InspectionDate", "DistressType", "Severity"]
        size_agg_func = "sum" if aggregate_sizes == "sum" else ("max" if aggregate_sizes == "max" else "first")

        aggregated_df = (
            main_df.groupby(group_keys, as_index=False)
            .agg({
                "InspectionUnit.": "first",
                "NoDistresses": "first",
                "Special": "first",
                "RawSize": size_agg_func,
                "InspectionArea": "sum",
                "InspectionLength": "sum",
                "Comment": "first",
            })
        )

        def _apply_section_totals(row):
            key_sec = (str(row["StreetID"]), str(row["SectionID"]), row["InspectionDate"])
            totals = section_totals.get(key_sec)
            if totals:
                row["InspectionLength"] = round1(totals["length"])
                row["InspectionArea"] = round1(totals["area"])
            return row

        aggregated_df = aggregated_df.apply(_apply_section_totals, axis=1)

        aggregated_df["DistressSize"] = aggregated_df["RawSize"].round().astype(int)

        # NEW: final safety filter (in case anything slips through)
        aggregated_df = aggregated_df[aggregated_df["DistressSize"] >= 1].copy()

        aggregated_df = aggregated_df.drop(columns=["RawSize"])
        aggregated_df["InspectionArea"] = aggregated_df["InspectionArea"].round().astype(int)
        aggregated_df["InspectionLength"] = aggregated_df["InspectionLength"].round().astype(int)

        main_df = aggregated_df

    # ----- outputs -----
    stable_main_csv = os.path.join("output", "InspectionReport.csv")
    stable_points_csv = os.path.join("output", "InspectionReport_points.csv")
    stable_points_gpkg = os.path.join("output", "InspectionReport_points.gpkg")
    stable_intersections_debug_csv = os.path.join("output", "InspectionReport_intersections_debug.csv")

    inspection_dates = list(main_df["InspectionDate"]) if "InspectionDate" in main_df.columns else []
    date_token = pick_date_token(frames_csv_path, inspection_dates)
    dated_dir = os.path.join("output", dt.date.today().isoformat())
    dated_main_csv = os.path.join(dated_dir, f"InspectionReport-{date_token}.csv")
    dated_points_csv = os.path.join(dated_dir, f"InspectionReport-{date_token}_points.csv")
    dated_points_gpkg = os.path.join(dated_dir, f"InspectionReport-{date_token}_points.gpkg")
    dated_excel = os.path.join(dated_dir, f"InspectionReport-{date_token}.xlsx")
    dated_intersections_debug_csv = os.path.join(dated_dir, f"InspectionReport-{date_token}_intersections_debug.csv")
    updated_input_csv = os.path.join(dated_dir, f"Frames_updated_lengths_distress-{date_token}.csv")

    for path in [
        stable_main_csv, stable_points_csv, stable_points_gpkg,
        stable_intersections_debug_csv,
        dated_main_csv, dated_points_csv, dated_points_gpkg,
        dated_excel, dated_intersections_debug_csv,
        updated_input_csv
    ]:
        ensure_dir(path)

    # write updated input CSV
    if updated_frames_df is not None and not updated_frames_df.empty:
        updated_frames_df.to_csv(updated_input_csv, index=False)
    else:
        # still write a header-only file so pipeline is predictable
        frames_df.head(0).to_csv(updated_input_csv, index=False)

    output_cols = [
        "StreetID", "SectionID", "InspectionUnit#", "InspectionDate",
        "InspectionArea", "InspectionLength",
        "DistressType", "Severity", "DistressSize",
        "NoDistresses", "Special", "Comments"
    ]

    if main_df.empty:
        empty_out = pd.DataFrame(columns=output_cols)
        empty_out.to_csv(stable_main_csv, index=False)
        empty_out.to_csv(dated_main_csv, index=False)
        write_excel(empty_out, dated_excel)
    else:
        output_df = main_df.rename(columns={
            "InspectionUnit.": "InspectionUnit#",
            "Comment": "Comments",
        }).copy()

        if "VideoName" in output_df.columns:
            output_df = output_df.drop(columns=["VideoName"])

        for col in output_cols:
            if col not in output_df.columns:
                output_df[col] = "" if col in ("Comments", "DistressType", "Severity", "StreetID", "SectionID", "InspectionDate", "NoDistresses", "Special") else 0

        # (already filtered DistressSize >= 1, but keep as safety)
        if "DistressSize" in output_df.columns:
            output_df = output_df[output_df["DistressSize"].astype(int) >= 1].copy()

        output_df[output_cols].to_csv(stable_main_csv, index=False)
        output_df[output_cols].to_csv(dated_main_csv, index=False)
        write_excel(output_df[output_cols], dated_excel)

    # points outputs
    points_df = pd.DataFrame(point_rows_output, columns=[
        "PointID", "X_EPSG3857", "Y_EPSG3857", "PolylineID", "RoadName", "Status",
        "R1_ft", "R2_ft"
    ])

    if not points_df.empty:
        points_df["VideoName"] = points_df["PointID"].apply(extract_video_name)
        unique_videos = [v for v in sorted(points_df["VideoName"].unique()) if v]
        video_to_id = {v: i + 1 for i, v in enumerate(unique_videos)}
        points_df["VideoGroupID"] = points_df["VideoName"].map(video_to_id).fillna(0).astype(int)
    else:
        points_df["VideoName"] = []
        points_df["VideoGroupID"] = []

    points_cols = [
        "PointID", "X_EPSG3857", "Y_EPSG3857",
        "PolylineID", "RoadName", "Status",
        "R1_ft", "R2_ft", "VideoGroupID", "VideoName"
    ]
    points_df[points_cols].to_csv(stable_points_csv, index=False, float_format="%.2f")
    points_df[points_cols].to_csv(dated_points_csv, index=False, float_format="%.2f")

    points_gpkg_written = "none"
    if not points_df.empty:
        points_gdf = gpd.GeoDataFrame(
            points_df.copy(),
            geometry=gpd.points_from_xy(points_df["X_EPSG3857"], points_df["Y_EPSG3857"]),
            crs="EPSG:3857"
        )
        try:
            safe_remove(stable_points_gpkg)
            points_gdf.to_file(stable_points_gpkg, layer="points", driver="GPKG")
            safe_remove(dated_points_gpkg)
            points_gdf.to_file(dated_points_gpkg, layer="points", driver="GPKG")
            points_gpkg_written = "stable+dated"
        except Exception as e:
            points_gpkg_written = f"failed ({e})"
    else:
        points_gpkg_written = "skipped (no points)"

    intersections_debug_written = "none"
    if intersection_debug_rows:
        debug_df = pd.DataFrame(intersection_debug_rows)
        debug_df.to_csv(stable_intersections_debug_csv, index=False)
        debug_df.to_csv(dated_intersections_debug_csv, index=False)
        intersections_debug_written = "stable+dated"

    intersections_count = 0
    unmatched_count = 0
    if not points_df.empty:
        intersections_count = int((points_df["Status"] == "Intersection").sum())
        unmatched_count = int((points_df["Status"] == "Not found").sum())

    summary = (
        f"Stable MAIN CSV:    {stable_main_csv}\n"
        f"Stable POINTS CSV:  {stable_points_csv} ({len(points_df)} rows)\n"
        f"Stable POINTS GPKG: {stable_points_gpkg} [{points_gpkg_written}]\n"
        f"Stable INTERSECTION DEBUG CSV: {stable_intersections_debug_csv} [{intersections_debug_written}]\n"
        f"Dated MAIN CSV:     {dated_main_csv}\n"
        f"Dated POINTS CSV:   {dated_points_csv}\n"
        f"Dated POINTS GPKG:  {dated_points_gpkg}\n"
        f"Dated INTERSECTION DEBUG CSV: {dated_intersections_debug_csv} [{intersections_debug_written}]\n"
        f"Dated Excel:        {dated_excel}\n"
        f"Updated INPUT CSV:  {updated_input_csv}\n"
        f"Intersections (points): {intersections_count}, Unmatched (points): {unmatched_count}\n"
        f"Include intersections in MAIN: {'YES' if include_intersections else 'NO'}\n"
        f"Use GPS-based length correction: {'YES' if use_gps_based_length else 'NO'}\n"
        f"Dropped distress rows where scaled size < {MIN_OUTPUT_DISTRESS_SIZE}"
    )
    return summary


# ---------- tiny UI helpers ----------
def pick_file(var: tk.StringVar, title: str, patterns):
    path = filedialog.askopenfilename(title=title, filetypes=patterns)
    if path:
        var.set(path)


def append_log(widget: ScrolledText, text: str):
    widget.configure(state="normal")
    widget.insert(tk.END, text if text.endswith("\n") else (text + "\n"))
    widget.see(tk.END)
    widget.configure(state="disabled")


# ---------- button handlers ----------
def on_run(roads_var, frames_var, include_intersections_var, use_gps_length_var, run_btn, log_box):
    roads_path = roads_var.get().strip()
    frames_csv_path = frames_var.get().strip()
    include_intersections = bool(include_intersections_var.get())
    use_gps_length = bool(use_gps_length_var.get())

    if not roads_path or not os.path.isfile(roads_path):
        messagebox.showerror("Missing input", "Select a valid roads shapefile (.shp).")
        return
    if not frames_csv_path or not os.path.isfile(frames_csv_path):
        messagebox.showerror("Missing input", "Select a valid frames CSV.")
        return

    run_btn.config(state="disabled", text="Running…")
    append_log(
        log_box,
        f"Running...\nRoads:  {roads_path}\nFrames: {frames_csv_path}\nInclude intersections: {include_intersections}\nUse GPS lengths: {use_gps_length}\n"
    )

    def work():
        try:
            summary = run_pipeline(
                roads_path,
                frames_csv_path,
                include_intersections=include_intersections,
                use_gps_based_length=bool(use_gps_length)
            )
            append_log(log_box, summary + "\n")
            messagebox.showinfo("Done", summary)
        except Exception as e:
            append_log(log_box, f"[ERROR] {e}\n")
            messagebox.showerror("Error", str(e))
        finally:
            run_btn.config(state="normal", text="Run")

    threading.Thread(target=work, daemon=True).start()


def on_sanitize(frames_var, qgis_py_var, qgis_proj_var, sanitize_btn, log_box):
    frames_csv_path = frames_var.get().strip()
    qgis_py = qgis_py_var.get().strip()
    qgis_project = qgis_proj_var.get().strip()
    query_script = DEFAULT_QUERY_SCRIPT

    if not os.path.isfile(qgis_py):
        messagebox.showerror("Missing QGIS Python", f"Cannot find python-qgis-ltr.bat:\n{qgis_py}")
        return
    if not os.path.isfile(query_script):
        messagebox.showerror("Missing query.py", f"Cannot find sanitizer script:\n{query_script}")
        return
    if not os.path.isfile(qgis_project):
        messagebox.showerror("Missing project", f"Cannot find QGIS project:\n{qgis_project}")
        return
    if not frames_csv_path or not os.path.isfile(frames_csv_path):
        messagebox.showerror("Missing frames CSV", "Select a valid frames CSV.")
        return

    append_log(log_box, "[Sanitizer] Launching QGIS Python to remove rows by Error layer PointID…")
    append_log(log_box, f"[Sanitizer] QGIS Python: {qgis_py}")
    append_log(log_box, f"[Sanitizer] Project:     {qgis_project}")
    append_log(log_box, f"[Sanitizer] Input CSV:   {frames_csv_path}")
    append_log(log_box, f"[Sanitizer] Script:      {query_script}")

    cmd = [qgis_py, query_script, frames_csv_path]

    sanitize_btn.config(state="disabled", text="Sanitizing…")
    run_subprocess_and_stream(cmd_list=cmd, cwd=SCRIPT_DIR, log_box=log_box, btn=sanitize_btn, done_label="Sanitize")


# ---------- GUI ----------
root = tk.Tk()
root.title("Inspection Report – GUI")
root.geometry("980x580")

roads_var = tk.StringVar()
frames_var = tk.StringVar()
qgis_py_var = tk.StringVar(value=DEFAULT_QGIS_PY_BAT)
qgis_proj_var = tk.StringVar(value=DEFAULT_QGIS_PROJECT)
include_intersections_var = tk.BooleanVar(value=False)
use_gps_length_var = tk.BooleanVar(value=DEFAULT_GPS_BASED_LENGTH)

pad = {"padx": 8, "pady": 6}

# Row 0: Roads
tk.Label(root, text="Surveyed Roads shapefile (.shp):").grid(row=0, column=0, sticky="w", **pad)
tk.Entry(root, textvariable=roads_var, width=78).grid(row=0, column=1, sticky="we", **pad)
tk.Button(root, text="Browse", command=lambda: pick_file(
    roads_var, "Select roads shapefile (.shp)", [("Shapefile", "*.shp"), ("All files", "*.*")]
)).grid(row=0, column=2, **pad)

# Row 1: Frames
tk.Label(root, text="Input CSV:").grid(row=1, column=0, sticky="w", **pad)
tk.Entry(root, textvariable=frames_var, width=78).grid(row=1, column=1, sticky="we", **pad)
tk.Button(root, text="Browse", command=lambda: pick_file(
    frames_var, "Select frames CSV", [("CSV files", "*.csv"), ("All files", "*.*")]
)).grid(row=1, column=2, **pad)

# Row 2: QGIS Python
tk.Label(root, text="QGIS Python (python-qgis-ltr.bat):").grid(row=2, column=0, sticky="w", **pad)
tk.Entry(root, textvariable=qgis_py_var, width=78).grid(row=2, column=1, sticky="we", **pad)
tk.Button(root, text="Browse", command=lambda: pick_file(
    qgis_py_var, "Select python-qgis-ltr.bat", [("Batch files", "*.bat"), ("All files", "*.*")]
)).grid(row=2, column=2, **pad)

# Row 3: QGIS Project
tk.Label(root, text="QGIS Project (.qgz):").grid(row=3, column=0, sticky="w", **pad)
tk.Entry(root, textvariable=qgis_proj_var, width=78).grid(row=3, column=1, sticky="we", **pad)
tk.Button(root, text="Browse", command=lambda: pick_file(
    qgis_proj_var, "Select QGIS project", [("QGIS Project", "*.qgz"), ("All files", "*.*")]
)).grid(row=3, column=2, **pad)

# Row 4: Options + Action Buttons
include_chk = tk.Checkbutton(
    root,
    text="Include intersections in MAIN output (assign to main road)",
    variable=include_intersections_var
)
include_chk.grid(row=4, column=1, sticky="w", **pad)

gps_chk = tk.Checkbutton(
    root,
    text="Use GPS-based frame lengths",
    variable=use_gps_length_var
)
gps_chk.grid(row=4, column=0, sticky="w", **pad)

run_btn = tk.Button(
    root, text="Run", width=22,
    command=lambda: on_run(
        roads_var, frames_var, include_intersections_var,
        use_gps_length_var, run_btn, log_box
    )
)
run_btn.grid(row=4, column=1, sticky="e", **pad)

sanitize_btn = tk.Button(
    root, text="Sanitize (remove Error IDs)", width=26,
    command=lambda: on_sanitize(frames_var, qgis_py_var, qgis_proj_var, sanitize_btn, log_box)
)
sanitize_btn.grid(row=4, column=2, **pad)

# Row 5: Log
log_box = ScrolledText(root, height=18, state="disabled")
tk.Label(root, text="Log:").grid(row=5, column=0, sticky="nw", **pad)
log_box.grid(row=5, column=1, columnspan=2, sticky="nsew", padx=8, pady=(0, 8))

# Clear log button
clear_btn = tk.Button(
    root, text="Clear Log", width=16,
    command=lambda: (log_box.configure(state="normal"),
                     log_box.delete("1.0", tk.END),
                     log_box.configure(state="disabled"))
)
clear_btn.grid(row=6, column=2, sticky="e", **pad)

root.grid_columnconfigure(1, weight=1)
root.grid_rowconfigure(5, weight=1)

root.mainloop()
